"""Regulatory burden estimate, following the OIA Regulatory Burden Estimation Tool.

Method (Regulatory Burden Estimation Tool, July 2026, sheets 1 and 3):
- Three cost categories: administrative compliance (staff time), substantive
  compliance (purchased services and products) and delay costs.
- One-off costs fall in year 1. Recurring costs repeat each year; by default from
  year 1 (the tool's "Initial Cost (Year 1)" includes the first year of recurring
  costs), or from year 2 when a recurring task only starts after a one-off set-up.
- The estimate covers the period of analysis (default 10 years), undiscounted.
  The 10-year total is compared with IA threshold 1 ($20 million); the annual
  average is the point estimate reported in the IA.

Costs are signed: a negative figure is a reduction in burden (a saving).
"""
from __future__ import annotations

import io
from typing import Literal

from pydantic import BaseModel, Field

from core.framework import load_framework

THRESHOLD_1 = 20_000_000  # IA threshold 1: $20 million in total over 10 years

CostType = Literal["Work-related", "Non-work-related"]
CATEGORY_LABELS = {
    "administrative": "Administrative Compliance Costs",
    "substantive": "Substantive Compliance Costs",
    "delay": "Delay Costs",
}


def _defaults() -> dict[str, float]:
    rbe = load_framework().rbe_defaults
    return {s.id: s.default for s in rbe.general_settings if s.default is not None}


class Settings(BaseModel):
    entities: float = Field(0, description="Number of entities affected (required)")
    people_per_entity: float = Field(default_factory=lambda: _defaults()["people_per_entity"])
    work_labour_rate: float = Field(default_factory=lambda: _defaults()["work_labour_rate"])
    non_work_labour_rate: float = Field(default_factory=lambda: _defaults()["non_work_labour_rate"])
    period_years: int = Field(default_factory=lambda: int(_defaults()["period_years"]))


class AdminItem(BaseModel):
    description: str = ""
    setup_hours: float = 0  # one-off, per person
    ongoing_hours: float = 0  # per occasion, per person
    frequency: float = 0  # occasions per year
    cost_type: CostType = "Work-related"
    entities: float | None = None  # overrides the general setting
    people_per_entity: float | None = None
    recurring_from_year: int = 1


class SubstantiveItem(BaseModel):
    description: str = ""
    one_off_cost: float = 0  # per entity
    recurring_cost: float = 0  # per entity, per occasion
    frequency: float = 0  # times per year
    entities: float | None = None
    recurring_from_year: int = 1


class DelayItem(BaseModel):
    description: str = ""
    income_lost_per_day: float = 0  # standby expenses + lost income, per entity per day
    delay_days: float = 0  # length (or change in length) of the delay
    recurrence: float = 1  # delays per entity per year
    one_off: bool = False
    entities: float | None = None


class BurdenInputs(BaseModel):
    title: str = ""
    settings: Settings = Field(default_factory=Settings)
    administrative: list[AdminItem] = Field(default_factory=list)
    substantive: list[SubstantiveItem] = Field(default_factory=list)
    delay: list[DelayItem] = Field(default_factory=list)
    assumptions: str = ""


class ItemResult(BaseModel):
    category: str
    description: str
    years: list[float]

    @property
    def total(self) -> float:
        return sum(self.years)


class BurdenResult(BaseModel):
    period_years: int
    items: list[ItemResult]

    def category_years(self, category: str) -> list[float]:
        years = [0.0] * self.period_years
        for it in self.items:
            if it.category == category:
                years = [a + b for a, b in zip(years, it.years)]
        return years

    @property
    def years(self) -> list[float]:
        return [sum(v) for v in zip(*(self.category_years(c) for c in CATEGORY_LABELS))]

    @property
    def total(self) -> float:
        return sum(self.years)

    @property
    def annual_average(self) -> float:
        return self.total / self.period_years if self.period_years else 0.0

    @property
    def meets_threshold_1(self) -> bool:
        return abs(self.total) >= THRESHOLD_1


def _schedule(period: int, one_off: float, recurring: float, from_year: int) -> list[float]:
    years = [0.0] * period
    if period:
        years[0] += one_off
    for y in range(max(from_year, 1) - 1, period):
        years[y] += recurring
    return years


def calculate(inputs: BurdenInputs) -> BurdenResult:
    s = inputs.settings
    period = max(int(s.period_years), 1)
    items: list[ItemResult] = []
    for a in inputs.administrative:
        rate = s.work_labour_rate if a.cost_type == "Work-related" else s.non_work_labour_rate
        n = a.entities if a.entities is not None else s.entities
        people = a.people_per_entity if a.people_per_entity is not None else s.people_per_entity
        per_hour = rate * n * people
        items.append(ItemResult(
            category="administrative", description=a.description,
            years=_schedule(period, a.setup_hours * per_hour,
                            a.ongoing_hours * a.frequency * per_hour, a.recurring_from_year),
        ))
    for b in inputs.substantive:
        n = b.entities if b.entities is not None else s.entities
        items.append(ItemResult(
            category="substantive", description=b.description,
            years=_schedule(period, b.one_off_cost * n,
                            b.recurring_cost * b.frequency * n, b.recurring_from_year),
        ))
    for d in inputs.delay:
        n = d.entities if d.entities is not None else s.entities
        cost = d.income_lost_per_day * d.delay_days * d.recurrence * n
        years = _schedule(period, cost, 0, period + 1) if d.one_off else _schedule(period, 0, cost, 1)
        items.append(ItemResult(category="delay", description=d.description, years=years))
    return BurdenResult(period_years=period, items=items)


def fmt_money(value: float) -> str:
    sign = "-" if value < 0 else ""
    v = abs(value)
    if v >= 1_000_000:
        return f"{sign}${v / 1_000_000:,.2f} million"
    return f"{sign}${v:,.0f}"


def summary_markdown(inputs: BurdenInputs, result: BurdenResult) -> str:
    s = inputs.settings
    lines = [
        f"# Regulatory burden estimate{': ' + inputs.title if inputs.title else ''}",
        "",
        "Prepared with IA Navigator using the method of the OIA Regulatory Burden Estimation "
        "Tool (July 2026). Figures are undiscounted. Negative figures are reductions in burden.",
        "",
        "## Result",
        "",
        f"- Total over {result.period_years} years: **{fmt_money(result.total)}**",
        f"- Average annual change in regulatory burden: **{fmt_money(result.annual_average)}**",
        f"- IA threshold 1 ($20 million in total over 10 years): "
        f"**{'likely met' if result.meets_threshold_1 else 'not met on these inputs'}**",
        "",
        "| Burden category | Year 1 | Each later year (average) | Total |",
        "|---|---|---|---|",
    ]
    for key, label in CATEGORY_LABELS.items():
        ys = result.category_years(key)
        later = sum(ys[1:]) / (len(ys) - 1) if len(ys) > 1 else 0.0
        lines.append(f"| {label} | {fmt_money(ys[0])} | {fmt_money(later)} | {fmt_money(sum(ys))} |")
    lines += [
        "",
        "## General settings",
        "",
        f"- Number of entities affected: {s.entities:,g}",
        f"- People per entity with new obligations: {s.people_per_entity:,g}",
        f"- Work-related labour rate: ${s.work_labour_rate:,.2f}/hr",
        f"- Non-work-related labour rate: ${s.non_work_labour_rate:,.2f}/hr",
        f"- Period of analysis: {s.period_years} years",
        "",
        "## Cost items",
        "",
    ]
    for it in result.items:
        lines.append(f"- {CATEGORY_LABELS[it.category]}: {it.description or '(no description)'} — "
                     f"{fmt_money(it.total)} over {result.period_years} years")
    if not result.items:
        lines.append("- None entered.")
    if inputs.assumptions.strip():
        lines += ["", "## Assumptions and data sources", "", inputs.assumptions.strip()]
    return "\n".join(lines) + "\n"


def to_xlsx(inputs: BurdenInputs, result: BurdenResult) -> bytes:
    """Workbook with the inputs, item-by-year calculations and an Estimates sheet
    laid out like sheet 3 of the OIA tool."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    bold = Font(bold=True)
    wb = Workbook()
    est = wb.active
    est.title = "Estimates"
    est.append(["REGULATORY BURDEN ESTIMATES", inputs.title or "<PROPOSAL TITLE>"])
    est.append([])
    est.append(["Burden Category", f"{result.period_years} Year Total"]
               + [f"YEAR {i}" for i in range(1, result.period_years + 1)])
    for key, label in CATEGORY_LABELS.items():
        ys = result.category_years(key)
        est.append([label, sum(ys)] + ys)
    est.append(["Total", result.total] + result.years)
    est.append([])
    est.append([f"{result.period_years} Year Total (For Dashboard, $M)", result.total / 1e6])
    est.append([f"Annual Average over {result.period_years} years ($M)", result.annual_average / 1e6])
    for row in (est[3], est[7]):
        for c in row:
            c.font = bold

    s = inputs.settings
    gs = wb.create_sheet("General settings")
    gs.append(["Setting", "Value"])
    for label, value in [
        ("Number of entities affected", s.entities),
        ("Number of people per entity with new obligations", s.people_per_entity),
        ("Work-related labour rate ($/hr)", s.work_labour_rate),
        ("Non-work related labour rate ($/hr)", s.non_work_labour_rate),
        ("Period of analysis (years)", s.period_years),
    ]:
        gs.append([label, value])

    calc = wb.create_sheet("Calculator")
    calc.append(["Category", "Description"] + [f"YEAR {i}" for i in range(1, result.period_years + 1)] + ["Total"])
    for it in result.items:
        calc.append([CATEGORY_LABELS[it.category], it.description] + it.years + [it.total])

    inp = wb.create_sheet("Inputs")
    for name, rows in [("administrative", inputs.administrative),
                       ("substantive", inputs.substantive), ("delay", inputs.delay)]:
        inp.append([CATEGORY_LABELS[name]])
        if rows:
            fields = list(type(rows[0]).model_fields)
            inp.append(fields)
            for r in rows:
                inp.append([getattr(r, f) for f in fields])
        inp.append([])

    notes = wb.create_sheet("Scratch Pad")
    notes.append(["Assumptions, calculations, methodologies and data sources"])
    for line in (inputs.assumptions or "").splitlines():
        notes.append([line])
    for ws in wb.worksheets:
        ws.column_dimensions["A"].width = 48
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class BurdenSuggestion(BaseModel):
    """Claude's first-cut cost items (structured output). Every number is an assumption."""

    entities: float = Field(description="Number of entities affected, best estimate")
    entities_basis: str = Field(description="Source or reasoning for the entity count")
    administrative: list[AdminItem]
    substantive: list[SubstantiveItem]
    delay: list[DelayItem]
    assumptions: str = Field(description="Basis for each number, one line per item, with sources")


def rows_to_items(model: type[BaseModel], rows: list[dict]) -> list:
    """Table rows from the editor to items, skipping blank rows and treating blanks as defaults."""
    import math

    items = []
    for row in rows:
        clean = {k: v for k, v in row.items()
                 if v is not None and not (isinstance(v, float) and math.isnan(v)) and v != ""}
        numbers = [v for k, v in clean.items() if isinstance(v, (int, float)) and k != "recurring_from_year"]
        if not clean.get("description") and not any(numbers):
            continue  # blank row
        items.append(model.model_validate(clean))
    return items
