"""Self-assessment against the OIA Impact Analysis Self-Assessment Tool (July 2026).

Each criterion is rated Not completed, Adequate practice or Best practice. Scoring
follows the workbook's summary table: a question shows an alert when fewer of its
criteria than `alert_below` are rated adequate or best practice, and the whole IA
shows an alert when fewer than `total_alert_below` are.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from enum import Enum
from typing import Literal

from pydantic import BaseModel, Field

from core.config import REFERENCE_DIR
from core.framework import Framework, load_framework


class Rating(str, Enum):
    NOT_COMPLETED = "not_completed"
    ADEQUATE = "adequate"
    BEST = "best"


RATING_LABELS = {
    Rating.NOT_COMPLETED: "Not completed",
    Rating.ADEQUATE: "Adequate practice",
    Rating.BEST: "Best practice",
}


class CriterionResponse(BaseModel):
    rating: Rating = Rating.NOT_COMPLETED
    note: str = ""


class Assessment(BaseModel):
    assessed_document: str = ""
    updated_at: str = ""
    responses: dict[str, CriterionResponse] = Field(default_factory=dict)

    def rating(self, criterion_id: str) -> Rating:
        r = self.responses.get(criterion_id)
        return r.rating if r else Rating.NOT_COMPLETED

    def touch(self) -> None:
        self.updated_at = datetime.now(timezone.utc).isoformat(timespec="seconds")


class QuestionScore(BaseModel):
    question: int
    label: str
    not_completed: int
    adequate: int
    best: int
    alert: str


class Score(BaseModel):
    questions: list[QuestionScore]
    total_not_completed: int
    total_adequate: int
    total_best: int
    total_alert: str

    @property
    def gaps(self) -> int:
        return self.total_not_completed


def score(assessment: Assessment, fw: Framework | None = None) -> Score:
    fw = fw or load_framework()
    sc = fw.self_assessment_scoring
    rows = []
    for q in fw.ia_questions:
        ratings = [assessment.rating(c.id) for c in fw.criteria_for(q.number)]
        adequate = ratings.count(Rating.ADEQUATE)
        best = ratings.count(Rating.BEST)
        rows.append(QuestionScore(
            question=q.number,
            label=sc.question_labels[q.number],
            not_completed=ratings.count(Rating.NOT_COMPLETED),
            adequate=adequate,
            best=best,
            alert=sc.question_alert_text if adequate + best < sc.alert_below[q.number] else "",
        ))
    adequate = sum(r.adequate for r in rows)
    best = sum(r.best for r in rows)
    return Score(
        questions=rows,
        total_not_completed=sum(r.not_completed for r in rows),
        total_adequate=adequate,
        total_best=best,
        total_alert=sc.total_alert_text if adequate + best < sc.total_alert_below else "",
    )


def report_markdown(assessment: Assessment, project_name: str, fw: Framework | None = None) -> str:
    fw = fw or load_framework()
    s = score(assessment, fw)
    lines = [
        f"# Self-assessment: {project_name}",
        "",
        "Prepared with IA Navigator against the OIA Impact Analysis Self-Assessment Tool "
        "(version July 2026). This is the agency's own assessment, prepared to support the "
        "Statement of Sufficiency. It is not an assessment by the Office of Impact Analysis.",
        "",
    ]
    if assessment.assessed_document:
        lines += [f"Document assessed: {assessment.assessed_document}", ""]
    if assessment.updated_at:
        lines += [f"Last updated: {assessment.updated_at}", ""]
    lines += [
        "## Summary",
        "",
        "| IA question | Not completed | Adequate practice | Best practice | Alert |",
        "|---|---|---|---|---|",
    ]
    for r in s.questions:
        lines.append(f"| {r.label} | {r.not_completed} | {r.adequate} | {r.best} | {r.alert} |")
    lines.append(f"| Totals | {s.total_not_completed} | {s.total_adequate} | {s.total_best} | {s.total_alert} |")
    lines += ["", "## Gaps to address", ""]
    gaps = [c for c in fw.self_assessment_criteria if assessment.rating(c.id) == Rating.NOT_COMPLETED]
    if not gaps:
        lines.append("No criteria are rated Not completed.")
    for c in gaps:
        note = assessment.responses.get(c.id, CriterionResponse()).note
        lines.append(f"- Q{c.question}: {c.criterion.splitlines()[0]}")
        lines.append(f"  - To reach adequate practice: {c.adequate_practice}")
        if note:
            lines.append(f"  - Note: {note}")
    lines += ["", "## All criteria", ""]
    for q in fw.ia_questions:
        lines.append(f"### Question {q.number}: {q.question}")
        lines.append("")
        for c in fw.criteria_for(q.number):
            resp = assessment.responses.get(c.id, CriterionResponse())
            lines.append(f"- {c.criterion.splitlines()[0]} — **{RATING_LABELS[resp.rating]}**")
            if resp.note:
                lines.append(f"  - {resp.note}")
        lines.append("")
    return "\n".join(lines)


def filled_workbook(assessment: Assessment, fw: Framework | None = None) -> bytes:
    """The OIA Self-Assessment Tool workbook with this assessment's answers filled in.

    Each criterion row has three TRUE/FALSE cells: B (not completed), D (adequate) and
    F (best practice). The workbook's own formulas then produce its summary table.
    """
    from openpyxl import load_workbook

    fw = fw or load_framework()
    path = REFERENCE_DIR / fw.meta.sources["self_assessment_tool"].local_file
    wb = load_workbook(path)
    ws = wb["2. Self-Assessment tool"]
    for c in fw.self_assessment_criteria:
        row = int(c.citation.location.split()[1])
        rating = assessment.rating(c.id)
        ws.cell(row, 2).value = rating == Rating.NOT_COMPLETED
        ws.cell(row, 4).value = rating == Rating.ADEQUATE
        ws.cell(row, 6).value = rating == Rating.BEST
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Claude-suggested ratings (structured output) ---------------------------------

class SuggestedRating(BaseModel):
    criterion_id: str
    rating: Literal["not_completed", "adequate", "best"]
    reason: str = Field(description="One or two sentences citing what the draft does or lacks")
    to_improve: str = Field(description="What to add or change; empty if best practice")


class SuggestedAssessment(BaseModel):
    ratings: list[SuggestedRating]
    overall: str = Field(description="Three or four sentences on the main gaps, most important first")


def criteria_prompt_text(fw: Framework | None = None) -> str:
    """The rubric as text, for the cached part of the system prompt."""
    fw = fw or load_framework()
    lines = ["Self-Assessment Tool rubric (OIA, version July 2026):", ""]
    for q in fw.ia_questions:
        lines.append(f"Question {q.number}: {q.question}")
        for c in fw.criteria_for(q.number):
            lines.append(f"- [{c.id}] {' '.join(c.criterion.split())}")
            lines.append(f"  Adequate practice: {c.adequate_practice}")
            lines.append(f"  Best practice: {c.best_practice}")
        lines.append("")
    return "\n".join(lines).strip()


def apply_suggestions(assessment: Assessment, suggested: SuggestedAssessment, fw: Framework | None = None) -> int:
    """Copy suggested ratings into the assessment (only known criterion ids). Returns count."""
    fw = fw or load_framework()
    known = {c.id for c in fw.self_assessment_criteria}
    n = 0
    for s in suggested.ratings:
        if s.criterion_id not in known:
            continue
        note = s.reason + (f" To improve: {s.to_improve}" if s.to_improve else "")
        assessment.responses[s.criterion_id] = CriterionResponse(rating=Rating(s.rating), note=note.strip())
        n += 1
    return n
