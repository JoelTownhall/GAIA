"""Draft and export: an IA draft structured on the 7 IA questions, exported to Word.

Section structure follows the Framework (section 2.1, p. 12; Figure 2, p. 13): a
Dashboard IA gives concise analysis against the 7 IA questions plus a high-level
discussion of strategic priorities; a Detailed IA adds in-depth analysis, including
distributional impacts, with a Dashboard IA as its executive summary. The OIA
publishes official templates on its website; agencies should move the final text into
the current template. Sections the Practical Guide marks for removal before
publication (p. 18) are flagged `confidential`.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Literal

from pydantic import BaseModel, Field

from core.burden import CATEGORY_LABELS, BurdenInputs, calculate, fmt_money
from core.claude_client import build_system
from core.framework import load_framework
from core.prompts import DRAFT_MODE, GUARDRAILS

IAType = Literal["dashboard", "detailed"]
IA_TYPE_LABELS = {"dashboard": "Dashboard IA", "detailed": "Detailed IA"}


class Section(BaseModel):
    id: str
    title: str
    question: int | None = None
    guidance: str = ""
    words: dict[str, int]
    confidential: bool = False
    detailed_only: bool = False


SECTIONS: list[Section] = [
    Section(id="executive_summary", title="Executive summary", detailed_only=True,
            guidance="A concise summary of the analysis across the 7 IA questions, in the "
                     "manner of a Dashboard IA (Framework, Figure 2, p. 13).",
            words={"dashboard": 0, "detailed": 600}),
    Section(id="q1", title="The policy problem", question=1, words={"dashboard": 220, "detailed": 900}),
    Section(id="q2", title="Policy objectives", question=2, words={"dashboard": 150, "detailed": 600}),
    Section(id="q3", title="Policy options", question=3, words={"dashboard": 200, "detailed": 900}),
    Section(id="q4", title="Likely net benefit of each option", question=4,
            words={"dashboard": 320, "detailed": 1500}),
    Section(id="distributional", title="Distributional impacts", detailed_only=True,
            guidance="How costs and benefits fall on different cohorts, including small "
                     "businesses (Framework, Figure 2, p. 13; IA threshold 3, p. 10).",
            words={"dashboard": 0, "detailed": 700}),
    Section(id="q5", title="Consultation", question=5, words={"dashboard": 180, "detailed": 800}),
    Section(id="q6", title="Preferred option and implementation", question=6,
            words={"dashboard": 200, "detailed": 800}),
    Section(id="q7", title="Monitoring and evaluation", question=7, words={"dashboard": 150, "detailed": 600}),
    Section(id="strategic", title="Strategic considerations", confidential=True,
            guidance="High-level analysis of the option(s) against the government's key strategic "
                     "priorities (Framework, section 2.1, p. 12). Remove before publication "
                     "(Practical Guide, p. 18).",
            words={"dashboard": 120, "detailed": 400}),
]


def sections_for(ia_type: IAType) -> list[Section]:
    return [s for s in SECTIONS if ia_type == "detailed" or not s.detailed_only]


def section_guidance(section: Section) -> str:
    if section.question is None:
        return section.guidance
    q = load_framework().ia_questions[section.question - 1]
    lines = [f"IA question {q.number}: {q.question}"]
    for p in q.prompts:
        lines.append(f"- {p.text}")
        lines += [f"  - {i}" for i in p.items]
    if q.note:
        lines.append(q.note)
    return "\n".join(lines)


class Draft(BaseModel):
    ia_type: IAType = "dashboard"
    title: str = ""
    sections: dict[str, str] = Field(default_factory=dict)
    updated_at: str = ""

    def touch(self) -> None:
        self.updated_at = datetime.now(timezone.utc).isoformat(timespec="seconds")

    def filename(self) -> str:
        return f"draft_{self.ia_type}.json"


def draft_system() -> list[dict]:
    return build_system(GUARDRAILS, DRAFT_MODE, cached_context=load_framework().as_prompt_text())


def section_request(section: Section, ia_type: IAType, project_context: str, current_text: str = "",
                    instructions: str = "") -> list[dict]:
    words = section.words[ia_type]
    task = [
        f"Draft the section \"{section.title}\" of a {IA_TYPE_LABELS[ia_type]}.",
        f"Word guide: about {words} words." if words else "",
        "", "What the section must cover:", section_guidance(section),
    ]
    if section.id == "q4":
        task += ["", "Include the regulatory burden estimate from the project folder (average annual "
                 "change over 10 years, undiscounted) if there is one; otherwise insert a gap marker."]
    if current_text.strip():
        task += ["", "Current draft of this section (revise it rather than starting again):", current_text.strip()]
    if instructions.strip():
        task += ["", "Officer's instructions for this draft:", instructions.strip()]
    return [{
        "role": "user",
        "content": [
            {"type": "text", "text": f"<project>\n{project_context}\n</project>",
             "cache_control": {"type": "ephemeral"}},
            {"type": "text", "text": "\n".join(t for t in task if t is not None)},
        ],
    }]


# --- Export -----------------------------------------------------------------------

def _burden_table_rows(burden: BurdenInputs) -> list[list[str]]:
    result = calculate(burden)
    rows = [["Burden category", "Total over period", "Average per year"]]
    for key, label in CATEGORY_LABELS.items():
        total = sum(result.category_years(key))
        rows.append([label, fmt_money(total), fmt_money(total / result.period_years)])
    rows.append(["Total", fmt_money(result.total), fmt_money(result.annual_average)])
    return rows


def export_markdown(draft: Draft, project_name: str, agency: str = "", publication: bool = False,
                    burden: BurdenInputs | None = None) -> str:
    lines = [f"# {draft.title or project_name}", "", f"{IA_TYPE_LABELS[draft.ia_type]}"
             + (f" — {agency}" if agency else ""), ""]
    for s in sections_for(draft.ia_type):
        if publication and s.confidential:
            continue
        lines += [f"## {s.title}" + (" (remove before publication)" if s.confidential else ""), ""]
        lines += [draft.sections.get(s.id, "").strip() or "[Not yet drafted]", ""]
        if s.id == "q4" and burden is not None:
            rows = _burden_table_rows(burden)
            lines += ["| " + " | ".join(rows[0]) + " |", "|---|---|---|"]
            lines += ["| " + " | ".join(r) + " |" for r in rows[1:]] + [""]
    lines += ["## Statement of Sufficiency", "",
              "[To be completed and endorsed by a Deputy Secretary (SES Band 3), equivalent or higher "
              "(Practical Guide, p. 13). Use the current OIA template.]", ""]
    return "\n".join(lines)


def _add_markdown(doc, text: str) -> None:
    """Minimal Markdown to Word: headings, bullets, numbered lists, bold, paragraphs."""
    import re

    for block in text.strip().splitlines():
        line = block.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^(#{3,6})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2), level=3)
            continue
        if re.match(r"^\s*[-*]\s+", line):
            para = doc.add_paragraph(style="List Bullet")
            line = re.sub(r"^\s*[-*]\s+", "", line)
        elif re.match(r"^\s*\d+[.)]\s+", line):
            para = doc.add_paragraph(style="List Number")
            line = re.sub(r"^\s*\d+[.)]\s+", "", line)
        else:
            para = doc.add_paragraph()
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", line)):
            if part:
                para.add_run(part).bold = i % 2 == 1


def export_docx(draft: Draft, project_name: str, agency: str = "", publication: bool = False,
                burden: BurdenInputs | None = None) -> bytes:
    """Word document with real heading styles and a header row on tables (accessibility)."""
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = docx.Document()
    doc.core_properties.title = draft.title or project_name
    doc.core_properties.subject = IA_TYPE_LABELS[draft.ia_type]
    doc.add_heading(draft.title or project_name, level=0)
    doc.add_paragraph(IA_TYPE_LABELS[draft.ia_type] + (f" — {agency}" if agency else ""))
    if not publication:
        doc.add_paragraph().add_run(
            "Draft prepared with IA Navigator. Check against the current OIA template before "
            "submission. Sections marked for removal must be taken out before publication."
        ).italic = True
    for s in sections_for(draft.ia_type):
        if publication and s.confidential:
            continue
        doc.add_heading(s.title + (" (remove before publication)" if s.confidential else ""), level=1)
        _add_markdown(doc, draft.sections.get(s.id, "") or "[Not yet drafted]")
        if s.id == "q4" and burden is not None:
            rows = _burden_table_rows(burden)
            doc.add_paragraph("Regulatory burden estimate (undiscounted)")
            table = doc.add_table(rows=len(rows), cols=3)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    table.cell(r, c).text = val
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
            for cell in table.rows[0].cells:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_heading("Statement of Sufficiency", level=1)
    doc.add_paragraph(
        "[To be completed and endorsed by a Deputy Secretary (SES Band 3), equivalent or higher "
        "(Practical Guide, p. 13). Use the current OIA template.]"
    )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
