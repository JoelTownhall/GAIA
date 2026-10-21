"""Core logic behind the General advice, Research and Draft and export tools."""
import io

import docx
import pytest

from core.advice import build_messages, focus_label, format_passages, transcript_markdown, user_turn_text
from core.burden import BurdenInputs, Settings, SubstantiveItem
from core.context import build_project_context
from core.drafting import Draft, export_docx, export_markdown, section_request, sections_for
from core.project import Project
from core.research import FOCUS_AREAS, brief_markdown, research_request
from core.retrieval import Passage, SearchResult
from core.storage import LocalStore


@pytest.fixture
def store_and_project(tmp_path):
    store = LocalStore(root=tmp_path)
    p = store.create_project(Project.new("Food labelling", agency="DHDA", policy_summary="Labels confuse."))
    return store, p


def test_project_context_includes_text_and_lists_omissions(store_and_project):
    store, p = store_and_project
    store.write_text(p.id, "research", "brief.md", "Evidence here.")
    store.write_file(p.id, "brief", "photo.png", b"\x89PNG")
    store.write_text(p.id, "burden", "burden_inputs.json", "{}")
    store.write_text(p.id, "drafts", "long.md", "x" * 50)
    d = docx.Document()
    d.add_paragraph("From a Word file.")
    buf = io.BytesIO()
    d.save(buf)
    store.write_file(p.id, "brief", "minute.docx", buf.getvalue())

    ctx = build_project_context(store, p, max_file_chars=20)
    assert "Project: Food labelling" in ctx.text and "Agency: DHDA" in ctx.text
    assert "Evidence here." in ctx.text and "From a Word file." in ctx.text
    assert "Only the first 20 of 50 characters" in ctx.text
    assert "00_brief/photo.png (file type not supported)" in ctx.omitted
    assert not any("burden_inputs.json" in x for x in ctx.included + ctx.omitted)  # JSON skipped
    only = build_project_context(store, p, folders=["research"])
    assert only.included == ["01_research/brief.md"]


def test_advice_messages_cache_project_context_on_first_turn():
    msgs = build_messages("CTX", [], "Q1")
    assert msgs[0]["content"][0]["cache_control"] == {"type": "ephemeral"}
    assert "CTX" in msgs[0]["content"][0]["text"]
    history = [{"role": "user", "content": "Q1"}, {"role": "assistant", "content": "A1"}]
    msgs = build_messages("CTX", history, "Q2")
    assert [m["role"] for m in msgs] == ["user", "assistant", "user"]
    assert msgs[0]["content"][1]["text"] == "Q1" and msgs[2]["content"] == "Q2"


def test_advice_turn_and_passages():
    r = SearchResult(Passage("practical_guide", "Practical Guide title", 13, "Step 3 > SoS", "Deputy Secretary"), 5.0)
    text = format_passages([r])
    assert 'source="Practical Guide" page="13"' in text and "Deputy Secretary" in text
    assert format_passages([]).startswith("No passages")
    turn = user_turn_text("Who signs?", focus_label(3, 7), text)
    assert "Step 3: Assess the IA" in turn and "IA question 7" in turn and turn.endswith("Who signs?")
    md = transcript_markdown("P", [{"role": "user", "content": "Who signs?", "focus": "Step 3"},
                                   {"role": "assistant", "content": "A Deputy Secretary."}])
    assert "## Question" in md and "*Focus: Step 3*" in md and "A Deputy Secretary." in md


def test_research_request_and_brief():
    req = research_request("Labels confuse.", ["magnitude", "costs", "unknown"], "How many brands?", "CTX")
    assert FOCUS_AREAS["costs"] in req and "How many brands?" in req and "<project>" in req
    md = brief_markdown("Food", "Labels confuse.", ["magnitude"], "## Summary\nText.",
                        [{"url": "https://abs.gov.au", "title": "ABS", "cited": True},
                         {"url": "https://x.org", "title": "X", "cited": False}])
    assert "[ABS](https://abs.gov.au)" in md and "(consulted, not cited)" in md


def test_sections_by_ia_type():
    dash = [s.id for s in sections_for("dashboard")]
    det = [s.id for s in sections_for("detailed")]
    assert dash == ["q1", "q2", "q3", "q4", "q5", "q6", "q7", "strategic"]
    assert det[0] == "executive_summary" and "distributional" in det


def test_section_request_includes_question_prompts():
    s = next(s for s in sections_for("dashboard") if s.id == "q3")
    msgs = section_request(s, "dashboard", "CTX", current_text="Old", instructions="Shorter")
    task = msgs[0]["content"][1]["text"]
    assert "What policy options are being considered?" in task
    assert "at least one non-regulatory alternative" in task
    assert "Old" in task and "Shorter" in task and "about 200 words" in task


def test_exports_mark_and_drop_confidential_sections():
    draft = Draft(title="Food labelling IA", sections={"q1": "The problem.\n\n- point **one**", "strategic": "Secret."})
    burden = BurdenInputs(settings=Settings(entities=10), substantive=[SubstantiveItem(one_off_cost=100)])
    md = export_markdown(draft, "P", burden=burden)
    assert "Strategic considerations (remove before publication)" in md and "| Total |" in md
    pub = export_markdown(draft, "P", publication=True)
    assert "Secret." not in pub and "Statement of Sufficiency" in pub

    doc = docx.Document(io.BytesIO(export_docx(draft, "P", agency="DHDA", burden=burden)))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "The policy problem" in headings and "Statement of Sufficiency" in headings
    assert any(p.style.name == "List Bullet" for p in doc.paragraphs)
    assert len(doc.tables) == 1 and doc.tables[0].cell(0, 0).text == "Burden category"
    pub_doc = docx.Document(io.BytesIO(export_docx(draft, "P", publication=True)))
    assert "Secret." not in "\n".join(p.text for p in pub_doc.paragraphs)
