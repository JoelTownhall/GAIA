import io

import pytest
from openpyxl import load_workbook

from core.burden import (
    AdminItem,
    BurdenInputs,
    DelayItem,
    Settings,
    SubstantiveItem,
    calculate,
    summary_markdown,
    to_xlsx,
)
from core.framework import load_framework


def expected(example_id):
    ex = {e.id: e for e in load_framework().rbe_defaults.worked_examples}
    return ex[example_id].ten_year_total


def test_defaults_come_from_the_rbe_tool():
    s = Settings()
    assert (s.work_labour_rate, s.non_work_labour_rate, s.people_per_entity, s.period_years) == (91.54, 41, 1, 10)


def test_example_a_reporting():
    # 110 wholesalers, 20 + 2 hours a year at $150/hour
    inputs = BurdenInputs(
        settings=Settings(entities=110, work_labour_rate=150),
        administrative=[AdminItem(description="Collect and report", ongoing_hours=22, frequency=1)],
    )
    assert calculate(inputs).total == pytest.approx(expected("example_a"))


def test_example_b_certification_with_refresher():
    # 15,000 staff: 4-hour course once, then a 1-hour refresher every 3 years
    inputs = BurdenInputs(
        settings=Settings(entities=15000),
        administrative=[AdminItem(description="Certification", setup_hours=4,
                                  ongoing_hours=1, frequency=1 / 3, recurring_from_year=2)],
    )
    assert calculate(inputs).total == pytest.approx(expected("example_b"))


def test_example_c_delay_reduction():
    # 9 approvals a year, 16 fewer days of delay at $13,800 a day: a saving
    inputs = BurdenInputs(
        settings=Settings(entities=9),
        delay=[DelayItem(description="Faster approvals", income_lost_per_day=13800,
                         delay_days=35 - 19, recurrence=1)],
    )
    assert calculate(inputs).total == pytest.approx(expected("example_c"))
    saving = BurdenInputs(settings=Settings(entities=9),
                          delay=[DelayItem(income_lost_per_day=13800, delay_days=-16)])
    assert calculate(saving).total == pytest.approx(-expected("example_c"))


def test_example_d_system_purchase_and_upkeep():
    inputs = BurdenInputs(
        settings=Settings(entities=30000),
        substantive=[SubstantiveItem(description="Buy and install system", one_off_cost=800)],
        administrative=[AdminItem(description="Set up and maintain", setup_hours=2,
                                  ongoing_hours=1, frequency=1, recurring_from_year=2)],
    )
    result = calculate(inputs)
    assert result.total == pytest.approx(expected("example_d"))
    assert result.meets_threshold_1
    assert result.annual_average == pytest.approx(expected("example_d") / 10)


def test_year_one_includes_first_year_of_recurring_costs():
    inputs = BurdenInputs(settings=Settings(entities=10),
                          substantive=[SubstantiveItem(one_off_cost=100, recurring_cost=50, frequency=2)])
    years = calculate(inputs).years
    assert years[0] == pytest.approx(10 * 100 + 10 * 50 * 2)
    assert years[1] == pytest.approx(10 * 50 * 2)
    assert len(years) == 10


def test_one_off_delay_and_item_overrides():
    inputs = BurdenInputs(
        settings=Settings(entities=100),
        delay=[DelayItem(income_lost_per_day=1000, delay_days=5, one_off=True, entities=2)],
        administrative=[AdminItem(ongoing_hours=1, frequency=1, cost_type="Non-work-related",
                                  people_per_entity=3)],
    )
    r = calculate(inputs)
    assert r.category_years("delay") == [10000] + [0] * 9
    assert r.category_years("administrative")[0] == pytest.approx(41 * 100 * 3)


def test_threshold_uses_absolute_change():
    small = calculate(BurdenInputs(settings=Settings(entities=1),
                                   substantive=[SubstantiveItem(one_off_cost=-19_000_000)]))
    big = calculate(BurdenInputs(settings=Settings(entities=1),
                                 substantive=[SubstantiveItem(one_off_cost=-20_000_000)]))
    assert not small.meets_threshold_1
    assert big.meets_threshold_1


def test_outputs_summary_and_workbook():
    inputs = BurdenInputs(title="Test", settings=Settings(entities=30000), assumptions="Counts from ABS.",
                          substantive=[SubstantiveItem(description="System", one_off_cost=800)])
    result = calculate(inputs)
    md = summary_markdown(inputs, result)
    assert "$24.00 million" in md and "likely met" in md and "Counts from ABS." in md
    wb = load_workbook(io.BytesIO(to_xlsx(inputs, result)))
    assert wb.sheetnames == ["Estimates", "General settings", "Calculator", "Inputs", "Scratch Pad"]
    est = wb["Estimates"]
    assert est["A9"].value.startswith("10 Year Total") and est["B9"].value == pytest.approx(24.0)
    assert est["B7"].value == pytest.approx(24_000_000)
