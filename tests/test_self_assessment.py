import io

from openpyxl import load_workbook

from core.framework import load_framework
from core.self_assessment import (
    Assessment,
    CriterionResponse,
    Rating,
    SuggestedAssessment,
    SuggestedRating,
    apply_suggestions,
    criteria_prompt_text,
    filled_workbook,
    report_markdown,
    score,
)

fw = load_framework()


def all_rated(rating):
    return Assessment(responses={c.id: CriterionResponse(rating=rating) for c in fw.self_assessment_criteria})


def test_empty_assessment_alerts_everywhere():
    s = score(Assessment())
    assert s.total_not_completed == 26
    assert all(q.alert for q in s.questions)
    assert s.total_alert


def test_all_adequate_clears_alerts():
    s = score(all_rated(Rating.ADEQUATE))
    assert s.total_adequate == 26 and s.total_not_completed == 0
    assert not any(q.alert for q in s.questions)
    assert not s.total_alert


def test_one_gap_alerts_its_question_but_not_the_total():
    # Matches the workbook: the total alert only fires below 25 of 26
    a = all_rated(Rating.BEST)
    a.responses["q4_3"] = CriterionResponse(rating=Rating.NOT_COMPLETED)
    s = score(a)
    alerts = {q.question: q.alert for q in s.questions}
    assert alerts[4] and not any(v for k, v in alerts.items() if k != 4)
    assert not s.total_alert
    a.responses["q1_1"] = CriterionResponse(rating=Rating.NOT_COMPLETED)
    assert score(a).total_alert


def test_report_lists_gaps():
    a = all_rated(Rating.ADEQUATE)
    a.responses["q3_1"] = CriterionResponse(rating=Rating.NOT_COMPLETED, note="No non-regulatory option.")
    md = report_markdown(a, "Food labelling")
    assert "Gaps to address" in md and "No non-regulatory option." in md
    assert "not an assessment by the Office of Impact Analysis" in md


def test_filled_workbook_sets_checkboxes():
    a = all_rated(Rating.ADEQUATE)
    a.responses["q1_1"] = CriterionResponse(rating=Rating.BEST)
    ws = load_workbook(io.BytesIO(filled_workbook(a)))["2. Self-Assessment tool"]
    assert (ws["B16"].value, ws["D16"].value, ws["F16"].value) == (False, False, True)
    assert (ws["B17"].value, ws["D17"].value, ws["F17"].value) == (False, True, False)
    assert ws["B5"].value == "=COUNTIF(B16:B21,TRUE)"  # workbook formulas kept


def test_apply_suggestions_ignores_unknown_ids():
    a = Assessment()
    n = apply_suggestions(a, SuggestedAssessment(overall="", ratings=[
        SuggestedRating(criterion_id="q2_1", rating="best", reason="Strong rationale.", to_improve=""),
        SuggestedRating(criterion_id="q9_9", rating="best", reason="", to_improve=""),
    ]))
    assert n == 1 and a.rating("q2_1") == Rating.BEST


def test_rubric_text_names_every_criterion():
    text = criteria_prompt_text()
    for c in fw.self_assessment_criteria:
        assert f"[{c.id}]" in text
