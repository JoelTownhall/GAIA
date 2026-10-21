import re

import pytest

from core.config import REFERENCE_DIR
from core.framework import Citation, load_framework


def _norm(text: str) -> str:
    text = text.replace("‑", "-").replace("–", "-").replace("−", "-")
    text = re.sub(r"\s+", " ", text).strip()
    return text.replace("/ ", "/")  # PDF line wraps after a slash


def _missing(fw, key):
    return not (REFERENCE_DIR / fw.meta.sources[key].local_file).exists()


@pytest.fixture(scope="module")
def fw():
    return load_framework()


@pytest.fixture(scope="module")
def framework_pages(fw):
    if _missing(fw, "framework"):
        pytest.skip("Framework PDF not in reference_docs/oia_2026/")
    import pdfplumber

    with pdfplumber.open(REFERENCE_DIR / fw.meta.sources["framework"].local_file) as pdf:
        return {i: _norm(p.extract_text() or "") for i, p in enumerate(pdf.pages, 1)}


@pytest.fixture(scope="module")
def sa_sheet(fw):
    if _missing(fw, "self_assessment_tool"):
        pytest.skip("Self-Assessment Tool not in reference_docs/oia_2026/")
    import openpyxl

    wb = openpyxl.load_workbook(REFERENCE_DIR / fw.meta.sources["self_assessment_tool"].local_file)
    return wb["2. Self-Assessment tool"]


def test_spec_loads_with_five_ordered_steps(fw):
    assert [s.number for s in fw.steps] == [1, 2, 3, 4, 5]
    assert fw.meta.effective_from == "2026-07-01"


def test_every_step_has_a_resolvable_source(fw):
    for step in fw.steps:
        src = fw.source_for(step)
        assert (step.source_url or src.url).startswith("https://oia.pmc.gov.au/")


def test_steps_are_verified_and_cited(fw):
    assert fw.is_verified
    for step in fw.steps:
        assert step.citations, step.id
        assert any(c.source == "framework" for c in step.citations), step.id


def test_prompt_text_is_deterministic_and_flags_unverified(fw):
    text = fw.as_prompt_text()
    assert text == fw.as_prompt_text()
    assert "Step 3: Assess the IA" in text
    assert "Question 4 - net benefit: What is the likely net benefit of the option?" in text
    if not fw.is_verified:
        assert "not yet verified" in text
    else:
        assert "not yet verified" not in text


def test_no_superseded_framework_language(fw):
    text = fw.as_prompt_text().lower()
    for stale in ("iala", "lord admiralty", "guide to policy impact analysis"):
        assert stale not in text


def test_every_citation_names_a_known_source_and_a_locator(fw):
    cits = fw.all_citations()
    assert len(cits) > 50
    for c in cits:
        assert c.source in fw.meta.sources
        assert c.section
        assert c.page is not None or c.location


def test_citation_requires_page_or_location():
    with pytest.raises(ValueError):
        Citation(source="framework", section="Step 1")


def test_seven_ia_questions_each_cited(fw):
    assert [q.number for q in fw.ia_questions] == list(range(1, 8))
    for q in fw.ia_questions:
        assert q.citation.source == "framework"
        assert q.citation.page in (14, 15, 16)
        assert q.prompts


def test_ia_questions_are_verbatim_on_cited_page(fw, framework_pages):
    for q in fw.ia_questions:
        page = framework_pages[q.citation.page]
        assert _norm(q.question) in page, q.id
        for p in q.prompts:
            assert _norm(p.text) in page, (q.id, p.text)
            for item in p.items:
                assert _norm(item) in page, (q.id, item)


def test_thresholds_match_framework(fw, framework_pages):
    assert [t.number for t in fw.ia_thresholds] == [1, 2, 3]
    page = framework_pages[10]
    for t in fw.ia_thresholds:
        # The Framework starts each threshold in lower case; the spec capitalises it.
        assert _norm(t.text).lower() in page.lower(), t.id


def test_self_assessment_criteria_cover_every_question(fw):
    assert len(fw.self_assessment_criteria) == 26
    for n in range(1, 8):
        crits = fw.criteria_for(n)
        assert crits, n
        assert fw.self_assessment_scoring.alert_below[n] == len(crits)
    ids = [c.id for c in fw.self_assessment_criteria]
    assert len(ids) == len(set(ids))


def test_every_criterion_has_a_citation(fw):
    for c in fw.self_assessment_criteria:
        assert c.citation.source == "self_assessment_tool"
        assert re.match(r"row \d+", c.citation.location), c.id


def test_self_assessment_criteria_are_verbatim(fw, sa_sheet):
    for c in fw.self_assessment_criteria:
        row = int(re.match(r"row (\d+)", c.citation.location).group(1))
        assert _norm(c.criterion) == _norm(sa_sheet.cell(row, 1).value), c.id
        assert _norm(c.adequate_practice) == _norm(sa_sheet.cell(row, 5).value), c.id
        assert _norm(c.best_practice) == _norm(sa_sheet.cell(row, 7).value), c.id


def test_criteria_match_framework_prompts(fw):
    prompts = {
        q.number: [_norm(p.text + " " + " ".join(p.items)).replace(" - ", " ") for p in q.prompts]
        for q in fw.ia_questions
    }
    for c in fw.self_assessment_criteria:
        crit = _norm(c.criterion.replace("\n- ", " ").replace("\n", " ")).replace(" - ", " ")
        crit = crit.replace("Building on IA question 1", "Build on IA question 1")
        assert crit in prompts[c.question], c.id


def test_rbe_defaults_match_workbook(fw):
    rbe = fw.rbe_defaults
    assert rbe.setting("work_labour_rate").default == 91.54
    assert rbe.setting("non_work_labour_rate").default == 41
    assert rbe.setting("people_per_entity").default == 1
    assert rbe.setting("period_years").default == 10
    assert rbe.setting("entities").required and rbe.setting("entities").default is None
    assert rbe.aggregation.horizon_years == 10
    assert rbe.aggregation.discounting == "none"
    if _missing(fw, "rbe_tool"):
        pytest.skip("RBE Tool not in reference_docs/oia_2026/")
    import openpyxl

    ws = openpyxl.load_workbook(REFERENCE_DIR / fw.meta.sources["rbe_tool"].local_file)["2. Calculator"]
    for s in rbe.general_settings:
        row = int(re.search(r"A(\d+)", s.citation.location).group(1))
        assert ws.cell(row, 1).value == s.label, s.id
        assert ws.cell(row, 4).value == s.guidance, s.id
        if s.default is not None:
            assert ws.cell(row, 3).value == s.default, s.id


def test_rbe_worked_examples_add_up(fw):
    for ex in fw.rbe_defaults.worked_examples:
        expr = re.search(r"\(=(.+)\)$", ex.answer).group(1)
        assert re.fullmatch(r"[\d.*+\-() ]+", expr)
        assert eval(expr) == pytest.approx(ex.ten_year_total), ex.id  # noqa: S307
        stated = float(re.match(r"\$([\d.]+) MILLION", ex.answer).group(1))
        assert round(ex.ten_year_total / 1e6, 2) == stated, ex.id


def test_special_cases_are_cited(fw):
    ids = {sc.id for sc in fw.special_cases}
    assert {"pm_exemption", "election_commitment", "sunsetting", "treaties"} <= ids
    for sc in fw.special_cases:
        assert sc.citations, sc.id


def test_source_files_present(fw):
    assert fw.missing_source_files() == []
